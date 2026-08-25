"""Turn a statement file into a plain table of string cells.

HTML, CSV and XLSX all reduce to `list[list[str]]`, so the statement parser only
has to understand one shape.
"""

from __future__ import annotations

import csv
import io
from html.parser import HTMLParser
from typing import Iterable

NBSP = " "


class _TableExtractor(HTMLParser):
    """Collects every <tr> in the document as a list of cell strings.

    MetaTrader statements are nested-table HTML written by the terminal, not by a
    templating engine, so tags are often unclosed; HTMLParser tolerates that where
    an XML parser would not.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        # Text outside any table cell — MetaTrader prints the account header there.
        self.preamble: list[str] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == "tr":
            self._flush_row()
            self._row = []
        elif tag in ("td", "th"):
            self._flush_cell()
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th"):
            self._flush_cell()
        elif tag == "tr":
            self._flush_row()

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)
        else:
            text = " ".join(data.replace(NBSP, " ").split())
            if text:
                self.preamble.append(text)

    def _flush_cell(self) -> None:
        if self._cell is not None and self._row is not None:
            text = "".join(self._cell).replace(NBSP, " ")
            self._row.append(" ".join(text.split()))
        self._cell = None

    def _flush_row(self) -> None:
        self._flush_cell()
        if self._row:
            self.rows.append(self._row)
        self._row = None

    def close(self) -> None:  # noqa: D102 - flush whatever the document left open
        super().close()
        self._flush_row()


def html_rows(payload: bytes | str) -> list[list[str]]:
    text = _decode(payload)
    parser = _TableExtractor()
    parser.feed(text)
    parser.close()
    rows = parser.rows
    if parser.preamble:
        # Kept as a leading pseudo-row so account sniffing sees it like any other text.
        rows = [parser.preamble] + rows
    return rows


def csv_rows(payload: bytes | str) -> list[list[str]]:
    text = _decode(payload)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [[cell.replace(NBSP, " ").strip() for cell in row] for row in reader if row]


def xlsx_rows(payload: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the install profile
        raise RuntimeError(
            "Для XLSX нужен пакет openpyxl (pip install openpyxl). "
            "Либо выгрузи отчёт из терминала в формате HTML."
        ) from exc

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    rows: list[list[str]] = []
    for sheet in workbook.worksheets:
        for raw in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value).strip() for value in raw]
            if any(cells):
                rows.append(cells)
    workbook.close()
    return rows


def read_rows(filename: str, payload: bytes) -> list[list[str]]:
    """Pick a reader from the file extension, sniffing the content as a fallback."""
    lowered = (filename or "").lower()
    if lowered.endswith((".html", ".htm")):
        return html_rows(payload)
    if lowered.endswith((".csv", ".tsv", ".txt")):
        return csv_rows(payload)
    if lowered.endswith((".xlsx", ".xlsm")):
        return xlsx_rows(payload)

    head = payload[:2048].lstrip()
    if head[:2] == b"PK":
        return xlsx_rows(payload)
    if b"<" in head[:200] and (b"<html" in head.lower() or b"<table" in head.lower()):
        return html_rows(payload)
    return csv_rows(payload)


def _decode(payload: bytes | str) -> str:
    if isinstance(payload, str):
        return payload
    # MT4 writes Windows-1251 for Russian builds and UTF-8/UTF-16 for others.
    for encoding in ("utf-8-sig", "utf-16", "cp1251", "latin-1"):
        try:
            return payload.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return payload.decode("utf-8", errors="replace")


def flatten(rows: Iterable[Iterable[str]]) -> list[list[str]]:
    return [list(row) for row in rows]
